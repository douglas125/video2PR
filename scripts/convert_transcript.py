#!/usr/bin/env python3
"""Parse external transcripts (Google Meet, MS Teams, Zoom) into canonical JSON format."""

import argparse
import html
import math
import re
import sys
from pathlib import Path

from diagnostics import json_bytes, write_outputs


def detect_format(input_path: Path) -> str:
    """Auto-detect transcript format from extension and content patterns."""
    ext = input_path.suffix.lower()

    if ext == ".docx":
        return "teams_docx"
    if ext == ".sbv":
        return "sbv"

    content = input_path.read_text(encoding="utf-8-sig")

    if ext == ".vtt" or content.strip().startswith("WEBVTT"):
        # Teams VTT has <v Speaker> tags
        if re.search(r"<v\s+[^>]+>", content):
            return "teams_vtt"
        # Zoom VTT has "Speaker Name: text" inline (after timestamp lines)
        # Look for pattern after a timestamp arrow line
        if re.search(
            r"-->\s*[\d:.]+.*\n[A-Z][\w\s]+:\s+\S",
            content,
        ):
            return "zoom_vtt"
        # Generic VTT fallback (no speaker detection)
        return "zoom_vtt"  # default VTT to zoom parser (handles both with/without speakers)

    # Google Meet .txt: lines like "Speaker Name (0:12:34)"
    if re.search(r"^.+ \(\d+:\d+:\d+\)$", content, re.MULTILINE):
        return "google_txt"

    # Zoom TXT: "Speaker Name   HH:MM:SS" with 2+ spaces, text on next line
    if re.search(r"^.+\s{2,}\d{2}:\d{2}:\d{2}\s*$", content, re.MULTILINE):
        return "zoom_txt"

    # SBV pattern: "0:00:01.234,0:00:05.678"
    if re.search(r"^\d+:\d+:\d+\.\d+,", content, re.MULTILINE):
        return "sbv"

    if ext == ".txt" and re.search(r"^[^:\n]+:\s+\S", content, re.MULTILINE):
        return "untimed_txt"
    raise ValueError(f"Could not detect transcript format for {input_path}")


def parse_sbv_timestamp(ts: str) -> float:
    return parse_vtt_timestamp(ts)


def parse_sbv(content: str) -> list[dict]:
    """Parse Google Meet SBV format."""
    segments = []
    blocks = re.split(r"\n\n+", content.strip())

    for block in blocks:
        lines = block.strip().split("\n")
        if len(lines) < 2:
            continue

        # First line: timestamp range
        ts_match = re.match(r"(\d+:\d+:\d+\.\d+),(\d+:\d+:\d+\.\d+)", lines[0])
        if not ts_match:
            continue

        start = parse_sbv_timestamp(ts_match.group(1))
        end = parse_sbv_timestamp(ts_match.group(2))
        text = "\n".join(lines[1:]).strip()

        # Check for speaker prefix like "Speaker Name: text"
        speaker = None
        speaker_match = re.match(r"^([^:]+):\s*(.+)$", text, re.DOTALL)
        if speaker_match:
            speaker = speaker_match.group(1).strip()
            text = speaker_match.group(2).strip()

        segments.append(
            {
                "start": start,
                "end": end,
                "text": text,
                "speaker": speaker,
                "words": [],
            }
        )

    return segments


def parse_vtt_timestamp(ts: str) -> float:
    """Parse a complete timestamp; never turn malformed input into time zero."""
    match = re.fullmatch(r"(?:(\d+):)?(\d{1,2}):(\d{2})\.(\d{1,3})", ts.strip())
    if not match:
        raise ValueError(f"Invalid transcript timestamp: {ts}")
    h, m, sec, fraction = match.groups()
    if int(m) >= 60 or int(sec) >= 60:
        raise ValueError(f"Invalid transcript timestamp: {ts}")
    return int(h or 0) * 3600 + int(m) * 60 + int(sec) + float("0." + fraction)


def split_speaker(text: str) -> tuple[str | None, str]:
    match = re.match(r"^([^:\n]+):[ \t]+(.+)$", text, re.DOTALL)
    labels = {"note", "time", "update", "action", "topic", "summary", "status"}
    if match:
        candidate = match.group(1).strip()
        if candidate and candidate[0].isupper() and candidate.casefold() not in labels:
            return candidate, match.group(2).strip()
    return None, text


def parse_vtt(content: str, inline_speakers: bool) -> list[dict]:
    """Parse cue blocks once; comments must begin a separate NOTE block."""
    content = content.lstrip("\ufeff").replace("\r\n", "\n").replace("\r", "\n")
    segments = []
    for block in re.split(r"\n[ \t]*\n+", content.strip()):
        lines = block.strip().splitlines()
        if not lines or re.match(r"^(?:WEBVTT|NOTE|STYLE|REGION)(?:[ \t]|$)", lines[0]):
            continue
        timestamp_index = next((i for i, line in enumerate(lines) if "-->" in line), None)
        if timestamp_index is None:
            print("Warning: skipped VTT block without timestamps", file=sys.stderr)
            continue
        match = re.fullmatch(r"(\S+)\s+-->\s+(\S+)(?:[ \t]+.*)?", lines[timestamp_index])
        try:
            if not match:
                raise ValueError("Malformed cue timing")
            start, end = (parse_vtt_timestamp(ts) for ts in match.groups())
            if end <= start:
                raise ValueError("Cue end must follow start")
        except ValueError as exc:
            print(f"Warning: skipped VTT cue: {exc}", file=sys.stderr)
            continue
        text = "\n".join(lines[timestamp_index + 1 :]).strip()
        speaker = None
        voice = re.match(r"<v(?:\.[^ >]+)*[ \t]+([^>]+)>", text)
        if voice:
            speaker = html.unescape(voice.group(1).strip())
        text = html.unescape(re.sub(r"<[^>]*>", "", text)).strip()
        if speaker is None and inline_speakers:
            speaker, text = split_speaker(text)
        if text:
            segments.append(
                {"start": start, "end": end, "text": text, "speaker": speaker, "words": []}
            )
    return segments


def parse_teams_vtt(content: str) -> list[dict]:
    return parse_vtt(content, inline_speakers=False)


def parse_zoom_vtt(content: str) -> list[dict]:
    return parse_vtt(content, inline_speakers=True)


def parse_untimed(content: str) -> list[dict]:
    """Preserve speaker-prefixed text with explicitly unknown timing."""
    segments = []
    for line in content.splitlines():
        if not line.strip():
            continue
        speaker, text = split_speaker(line.strip())
        if speaker is None and segments:
            segments[-1]["text"] += "\n" + text
        else:
            segments.append(
                {
                    "start": 0.0,
                    "end": 0.0,
                    "text": text,
                    "speaker": speaker,
                    "words": [],
                    "timing_source": "unknown",
                }
            )
    return segments


def parse_zoom_txt(content: str) -> list[dict]:
    """Parse Zoom TXT format with 'Speaker Name   HH:MM:SS' headers."""
    segments = []
    # Match lines like "Speaker Name   00:12:34" (2+ spaces between name and time)
    header_pattern = re.compile(r"^(.+?)\s{2,}(\d{2}:\d{2}:\d{2})\s*$", re.MULTILINE)
    headers = list(header_pattern.finditer(content))

    for i, match in enumerate(headers):
        speaker = match.group(1).strip()
        ts_str = match.group(2)
        parts = ts_str.split(":")
        start = int(parts[0]) * 3600 + int(parts[1]) * 60 + int(parts[2])

        # Text is between this header and the next
        text_start = match.end()
        text_end = headers[i + 1].start() if i + 1 < len(headers) else len(content)
        text = content[text_start:text_end].strip()

        # End time = next segment's start, or start + 30s for last segment
        end = (
            int(headers[i + 1].group(2).split(":")[0]) * 3600
            + int(headers[i + 1].group(2).split(":")[1]) * 60
            + int(headers[i + 1].group(2).split(":")[2])
            if i + 1 < len(headers)
            else start + 30
        )

        if text:
            segments.append(
                {
                    "start": float(start),
                    "end": float(end),
                    "text": text,
                    "speaker": speaker,
                    "words": [],
                }
            )

    return segments


def parse_google_txt(content: str) -> list[dict]:
    """Parse Google Meet .txt format with speaker headers like 'Name (H:MM:SS)'."""
    segments = []
    # Match lines like "Speaker Name (0:12:34)"
    header_pattern = re.compile(r"^(.+?)\s*\((\d+:\d+:\d+)\)\s*$", re.MULTILINE)
    headers = list(header_pattern.finditer(content))

    for i, match in enumerate(headers):
        speaker = match.group(1).strip()
        ts_str = match.group(2)
        parts = ts_str.split(":")
        start = int(parts[0]) * 3600 + int(parts[1]) * 60 + int(parts[2])

        # Text is between this header and the next
        text_start = match.end()
        text_end = headers[i + 1].start() if i + 1 < len(headers) else len(content)
        text = content[text_start:text_end].strip()

        # End time = next segment's start, or start + 30s for last segment
        end = (
            int(headers[i + 1].group(2).split(":")[0]) * 3600
            + int(headers[i + 1].group(2).split(":")[1]) * 60
            + int(headers[i + 1].group(2).split(":")[2])
            if i + 1 < len(headers)
            else start + 30
        )

        if text:
            segments.append(
                {
                    "start": float(start),
                    "end": float(end),
                    "text": text,
                    "speaker": speaker,
                    "words": [],
                }
            )

    return segments


def parse_teams_docx(input_path: Path) -> list[dict]:
    """Parse MS Teams .docx transcript."""
    from docx import Document

    doc = Document(str(input_path))
    segments = []
    current_speaker = None
    current_start = 0.0
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if not any(re.match(r"^(.+?)\s{2,}(\d+:\d+:\d+)\s*$", t) for t in paragraphs):
        return parse_untimed("\n".join(paragraphs))
    seen_header = False
    current_texts = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # Teams docx typically has speaker + timestamp lines followed by text
        # Pattern: "Speaker Name  0:12:34"
        header_match = re.match(r"^(.+?)\s{2,}(\d+:\d+:\d+)\s*$", text)
        if header_match:
            seen_header = True
            # Save previous segment
            if current_texts:
                segments.append(
                    {
                        "start": current_start,
                        "end": current_start + 30,  # placeholder until next segment
                        "text": " ".join(current_texts),
                        "speaker": current_speaker,
                        "words": [],
                    }
                )

            current_speaker = header_match.group(1).strip()
            ts_parts = header_match.group(2).split(":")
            current_start = float(
                int(ts_parts[0]) * 3600 + int(ts_parts[1]) * 60 + int(ts_parts[2])
            )
            current_texts = []
        elif seen_header:
            current_texts.append(text)

    # Save final segment
    if current_texts:
        segments.append(
            {
                "start": current_start,
                "end": current_start + 30,
                "text": " ".join(current_texts),
                "speaker": current_speaker,
                "words": [],
            }
        )

    # Fix end times: each segment ends when the next begins
    for i in range(len(segments) - 1):
        segments[i]["end"] = segments[i + 1]["start"]

    return segments


def convert(input_path: Path, output_dir: Path, fmt: str) -> None:
    """Convert external transcript to canonical format."""
    output_dir.mkdir(parents=True, exist_ok=True)

    # Detect format if auto
    if fmt == "auto":
        fmt = detect_format(input_path)

    print(f"Detected format: {fmt}")

    # Parse based on format
    if fmt == "teams_docx":
        segments = parse_teams_docx(input_path)
    else:
        content = input_path.read_text(encoding="utf-8-sig")
        parsers = {
            "sbv": parse_sbv,
            "teams_vtt": parse_teams_vtt,
            "zoom_vtt": parse_zoom_vtt,
            "zoom_txt": parse_zoom_txt,
            "google_txt": parse_google_txt,
            "untimed_txt": parse_untimed,
            "vtt": parse_zoom_vtt,  # generic VTT uses zoom parser (handles both)
        }
        parser = parsers.get(fmt)
        if parser is None:
            print(f"Unknown format: {fmt}", file=sys.stderr)
            sys.exit(1)
        segments = parser(content)

    if not segments:
        raise ValueError("No recognized transcript content; existing outputs were preserved")
    for seg in segments:
        if (
            not all(
                isinstance(seg[k], (int, float)) and math.isfinite(seg[k]) and seg[k] >= 0
                for k in ("start", "end")
            )
            or seg["end"] < seg["start"]
        ):
            raise ValueError("Transcript contains invalid or reversed timestamps")
        if "timing_source" not in seg:
            seg["timing_source"] = (
                "start_observed_end_estimated"
                if fmt in {"google_txt", "zoom_txt", "teams_docx"}
                else "observed"
            )
    print(f"Parsed {len(segments)} segments")

    has_speakers = any(seg["speaker"] for seg in segments)
    has_timestamps = all(seg["timing_source"] != "unknown" for seg in segments)

    format_labels = {
        "sbv": "google_meet_sbv",
        "teams_vtt": "ms_teams_vtt",
        "vtt": "vtt_generic",
        "zoom_vtt": "zoom_vtt",
        "zoom_txt": "zoom_txt",
        "google_txt": "google_meet_txt",
        "teams_docx": "ms_teams_docx",
    }
    meta = {
        "source": format_labels.get(fmt, fmt),
        "has_timestamps": has_timestamps,
        "has_speakers": has_speakers,
        "original_file": input_path.name,
        "segment_count": len(segments),
    }
    transcript = {"segments": segments}
    outputs = {
        "external_transcript.json": json_bytes(transcript),
        f"external_transcript_original{input_path.suffix}": input_path.read_bytes(),
        "external_transcript_meta.json": json_bytes(meta),
    }
    if has_timestamps:
        outputs["transcript.json"] = json_bytes(transcript)
    write_outputs(output_dir, outputs)
    print(f"Transcript and metadata saved to {output_dir}")


def main():
    parser = argparse.ArgumentParser(
        description="Convert external transcripts to canonical JSON format"
    )
    parser.add_argument("--input", required=True, help="Path to transcript file")
    parser.add_argument("--output-dir", required=True, help="Output directory")
    parser.add_argument(
        "--format",
        default="auto",
        choices=[
            "auto",
            "untimed_txt",
            "sbv",
            "vtt",
            "teams_vtt",
            "zoom_vtt",
            "zoom_txt",
            "google_txt",
            "teams_docx",
        ],
        help="Transcript format (default: auto-detect)",
    )
    args = parser.parse_args()

    input_path = Path(args.input).resolve()
    output_dir = Path(args.output_dir).resolve()

    if not input_path.exists():
        print(f"Input file not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    try:
        convert(input_path, output_dir, args.format)
    except (ValueError, OSError) as exc:
        parser.exit(1, f"Conversion failed: {exc}\n")


if __name__ == "__main__":
    main()
